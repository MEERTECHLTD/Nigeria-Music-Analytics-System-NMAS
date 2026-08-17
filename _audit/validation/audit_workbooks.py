#!/usr/bin/env python3
"""Structural audit of every Excel deliverable against the NBS presentation standard."""
import sys, csv
from pathlib import Path
from openpyxl import load_workbook
OUT = Path("NBS FINAL delivery/03_Excel_Deliveries")
findings=[]
def add(wb,sheet,issue,detail,sev):
    findings.append({"workbook":wb,"sheet":sheet,"issue":issue,"detail":detail,"severity":sev})
for p in sorted(OUT.glob("*.xlsx")):
    try:
        wb=load_workbook(p, read_only=False, data_only=False)
    except Exception as e:
        add(p.name,"-","file will not open",str(e)[:80],"CRITICAL"); continue
    names=wb.sheetnames
    if not any(s.lower().startswith("cover") or s.lower()=="metadata" for s in names):
        add(p.name,"-","no cover/metadata sheet","NBS standard requires provenance sheet","HIGH")
    for s in names:
        sh=wb[s]
        if sh.sheet_state!="visible":
            add(p.name,s,"hidden sheet","sheet_state=%s"%sh.sheet_state,"CRITICAL")
        if s.lower()=="cover": continue
        if sh.max_row is None or sh.max_row<2:
            add(p.name,s,"empty sheet","max_row=%s"%sh.max_row,"MEDIUM"); continue
        if sh.freeze_panes is None:
            add(p.name,s,"no freeze panes","header scrolls out of view","LOW")
        if not sh.auto_filter.ref:
            add(p.name,s,"no autofilter","NBS tables should be filterable","LOW")
        if sh.print_area in (None,""):
            add(p.name,s,"no print area","page layout undefined","LOW")
        # hidden rows/cols
        hidden_c=[k for k,v in sh.column_dimensions.items() if v.hidden]
        hidden_r=[k for k,v in sh.row_dimensions.items() if v.hidden]
        if hidden_c: add(p.name,s,"hidden columns",",".join(map(str,hidden_c))[:60],"CRITICAL")
        if hidden_r: add(p.name,s,"hidden rows",str(len(hidden_r))+" rows","CRITICAL")
        # number formats on numeric columns
        hdr=[c.value for c in sh[1]]
        generic=0; total=0
        for row in sh.iter_rows(min_row=2, max_row=min(sh.max_row,200)):
            for c in row:
                if isinstance(c.value,(int,float)):
                    total+=1
                    if c.number_format=="General": generic+=1
        if total and generic/total>0.9:
            add(p.name,s,"numbers unformatted","%d/%d numeric cells use General format"%(generic,total),"MEDIUM")
        # alignment of numerics
        misaligned=0
        for row in sh.iter_rows(min_row=2, max_row=min(sh.max_row,50)):
            for c in row:
                if isinstance(c.value,(int,float)) and (c.alignment.horizontal or "general")=="general":
                    misaligned+=1
        if misaligned>20:
            add(p.name,s,"numerals not right-aligned","%d cells default-aligned"%misaligned,"LOW")
    wb.close()
fields=["workbook","sheet","issue","detail","severity"]
with open("_audit/validation/workbook_audit.csv","w",newline="",encoding="utf-8") as h:
    w=csv.DictWriter(h,fieldnames=fields); w.writeheader(); w.writerows(findings)
import collections
print("workbooks audited: %d" % len(list(OUT.glob('*.xlsx'))))
print("findings: %d" % len(findings))
for sev in ("CRITICAL","HIGH","MEDIUM","LOW"):
    c=collections.Counter(f["issue"] for f in findings if f["severity"]==sev)
    if c:
        print("  %s:" % sev)
        for k,n in c.most_common(): print("    %-32s %d" % (k,n))
